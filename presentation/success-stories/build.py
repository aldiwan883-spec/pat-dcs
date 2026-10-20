"""PAT success-stories deck (16:9 PowerPoint, editable).

    python presentation/success-stories/build.py      # needs python-pptx, pillow, qrcode
    soffice --headless --convert-to pdf --outdir presentation/success-stories/out presentation/success-stories/out/*.pptx

Inputs: projects.json (exported from src/data/projects.ts by export-data.ts), stories.json (contractor / end user /
sector / location / photo per project), assets/generated (maps and logos from render.mjs), assets/photos (licensed
photos, see assets/photos/CREDITS.json), profile/img and public/images (PAT's own photos).
Fonts: Barlow and Barlow Condensed (presentation/fonts) — install them before opening the deck.
"""
import json
import os

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
GEN = os.path.join(HERE, "assets", "generated")
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)

NAVY, ORANGE, OTXT, ODARK = "102840", "E05826", "B8441A", "F07A4A"
BLUE, SAND, SLATE, LINE, WHITE = "2A7CAC", "F3F2EE", "3B5368", "D5DCE2", "FFFFFF"
H_XB, H_B, H_SB = "Barlow Condensed ExtraBold", "Barlow Condensed", "Barlow Condensed SemiBold"
B_R, B_M, B_SB = "Barlow", "Barlow Medium", "Barlow SemiBold"

projects = {p["no"]: p for p in json.load(open(os.path.join(HERE, "projects.json")))}
data = json.load(open(os.path.join(HERE, "stories.json")))
companies = data["companies"]
credits_path = os.path.join(HERE, "assets", "photos", "CREDITS.json")
credits = json.load(open(credits_path)) if os.path.exists(credits_path) else {}

COUNTRY = {"egypt": "Egypt", "saudi-arabia": "Saudi Arabia", "oman": "Oman"}
PLATFORM = {"ovation": "Emerson Ovation", "deltav": "Emerson DeltaV", "emerson": "Emerson DCS",
            "siemens": "Siemens S7 PLC", "alstom": "Alstom DCS", "other": ""}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


def period(p):
    return f"{p['start']}–{p['end']}" if p.get("end") and p["end"] != p["start"] else str(p["start"])


# ------------------------------------------------------------------ drawing helpers
def rect(slide, x, y, w, h, color, alpha=None, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    if alpha is not None:
        srgb = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
        srgb.append(a)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    return s


def text(slide, x, y, w, h, runs, size=12, font=B_R, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=None, line=1.0, upper=False, para_space=0):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if para_space and i:
            p.space_before = Pt(para_space)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t.upper() if (o.get("upper", upper)) else t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.color.rgb = rgb(o.get("color", color))
            f.bold = o.get("bold", False)
            sp = o.get("spacing", spacing)
            if sp:
                r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
    return tb


def picture_fill(slide, path, x, y, w, h):
    """Picture cropped to fill the box (like CSS object-fit: cover)."""
    iw, ih = Image.open(path).size
    box = w / h
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if iw / ih > box:
        excess = 1 - box / (iw / ih)
        pic.crop_left = pic.crop_right = excess / 2
    else:
        excess = 1 - (iw / ih) / box
        pic.crop_top = pic.crop_bottom = excess / 2
    return pic


def picture_fit(slide, path, x, y, w, h, align="center"):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    return slide.shapes.add_picture(path, Inches(px), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def brand(p):
    return os.path.join(ROOT, "brand", p)


def footer(slide, n, dark=False):
    col = "C9D3DC" if dark else SLATE
    picture_fit(slide, brand("png/logo/pat-symbol-reverse-256.png" if dark else "png/logo/pat-symbol-color-256.png"), 0.55, 6.98, 0.3, 0.3, align="left")
    text(slide, 0.95, 7.02, 6, 0.25, "POWER AUTOMATION TEAM  ·  SUCCESS STORIES", 8, H_SB, col, spacing=1.2)
    text(slide, SW - 1.55, 7.02, 1.0, 0.25, f"{n:02d}", 9, H_B, col, align=PP_ALIGN.RIGHT)


def eyebrow(slide, x, y, w, t, color=OTXT):
    text(slide, x, y, w, 0.3, t, 11, H_SB, color, spacing=2.2, upper=True)


def title(slide, x, y, w, h, t, size=34, color=NAVY):
    text(slide, x, y, w, h, t, size, H_XB, color, upper=True, line=0.92)


def photo_path(ref):
    if not ref:
        return None
    if ref.startswith("pat:"):
        return os.path.join(ROOT, "profile", "img", ref[4:] + ".jpg")
    if ref.startswith("public:"):
        return os.path.join(ROOT, "public", "images", ref[7:] + ".jpg")
    for p in (os.path.join(HERE, "assets", "photos", ref + ".jpg"), os.path.join(ROOT, "profile", "img", ref + ".jpg")):
        if os.path.exists(p):
            return p
    return None


def logo_path(key):
    return os.path.join(GEN, "logos", f"{key}.png")


slides_made = [0]


def new_slide():
    slides_made[0] += 1
    return prs.slides.add_slide(BLANK), slides_made[0]


# ------------------------------------------------------------------ slides
def cover():
    s, n = new_slide()
    picture_fill(s, os.path.join(ROOT, "profile", "img", "team-stack.jpg"), 0, 0, SW, SH)
    rect(s, 0, 0, SW, SH, NAVY, alpha=0.78)
    rect(s, 0, 0, 7.4, SH, NAVY, alpha=0.55)
    picture_fit(s, brand("png/logo/pat-weave-device-reverse-1600.png"), 7.2, 2.6, 7.6, 5.6)
    picture_fit(s, brand("png/logo/pat-lockup-horizontal-reverse-transparent-2000.png"), 0.7, 0.65, 3.6, 0.95, align="left")
    eyebrow(s, 0.75, 2.55, 6, "Company presentation · 2026", ODARK)
    title(s, 0.7, 2.9, 7.5, 1.9, "Success stories", size=72, color=WHITE)
    text(s, 0.75, 4.35, 6.4, 1.0, f"{len(projects)} reference projects in Egypt, Saudi Arabia and Oman — DCS commissioning, "
         "loop checks, start-up and training since 1991.", 16, B_M, "E5ECF2", line=1.15)
    text(s, 0.75, 6.75, 8, 0.3, "patdcs.com  ·  +20 109 991 2020  ·  info@patdcs.com", 11, H_SB, ODARK, spacing=1.5)
    notes(s, "Cover photo: PAT team on site (PAT photo).")


TRUSTED_CONTRACTORS = ["emerson", "hanwha", "daewoo", "ge", "orascom", "iberdrola", "kbr", "cmi", "skoda", "initec", "giza", "saudifal", "savola", "leoron"]
TRUSTED_OWNERS = ["aramco", "sec", "marafiq", "eehc", "phoenix", "nrea", "unitedsugar"]


def trusted_by():
    s, n = new_slide()
    eyebrow(s, 0.7, 0.55, 8, "Trusted by")
    title(s, 0.7, 0.85, 11, 0.8, "Contractors, vendors and plant owners")
    text(s, 0.7, 1.6, 11.5, 0.5, "Companies PAT engineers and technicians have delivered for since 1991.", 14, B_R, SLATE)

    def grid(keys, y0, label, cols=7, cell_w=1.62, cell_h=0.95, gap=0.12):
        text(s, 0.7, y0, 8, 0.3, label, 10.5, H_SB, OTXT, spacing=2, upper=True)
        for i, k in enumerate(keys):
            r_, c_ = divmod(i, cols)
            x, y = 0.7 + c_ * (cell_w + gap), y0 + 0.38 + r_ * (cell_h + gap)
            rect(s, x, y, cell_w, cell_h, WHITE, line=LINE)
            picture_fit(s, logo_path(k), x + 0.18, y + 0.2, cell_w - 0.36, cell_h - 0.4)
    grid(TRUSTED_CONTRACTORS, 2.25, "EPC contractors & system vendors")
    grid(TRUSTED_OWNERS, 4.95, "Plant owners & end users")
    footer(s, n)
    notes(s, "Logos are trademarks of their owners and are shown to identify companies PAT has worked for. "
             "Check each company's policy before external use — Saudi Aramco in particular requires written approval "
             "to reference its name or logo in marketing material. Name tiles are used where no logo file is available.")


def numbers():
    s, n = new_slide()
    picture_fill(s, os.path.join(ROOT, "profile", "img", "team-site.jpg"), 7.4, 0, SW - 7.4, SH)
    eyebrow(s, 0.7, 0.55, 6, "PAT in numbers")
    title(s, 0.7, 0.85, 6.4, 1.4, "The commissioning specialists")
    stats = [("1991", "Established in Egypt"), (str(len(projects)), "Reference projects"), ("3", "Countries: Egypt, Saudi Arabia, Oman"),
             ("12 GW+", "Combined capacity of the plants we worked on"), ("24/7", "Support for our clients")]
    for i, (v, l) in enumerate(stats):
        y = 2.3 + i * 0.88
        text(s, 0.7, y, 2.3, 0.8, v, 40, H_XB, ORANGE, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 3.05, y, 4.0, 0.8, l, 14, B_M, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        if i:
            rect(s, 0.7, y - 0.04, 6.3, 0.012, LINE)
    footer(s, n)
    notes(s, "12 GW+: sum of the plant capacities in the reference list, counting each plant once "
             "(Rabigh II and Shoaiba Stage I not counted). Photo: PAT team on site.")


def sectors():
    s, n = new_slide()
    eyebrow(s, 0.7, 0.55, 8, "Sectors we serve")
    title(s, 0.7, 0.85, 11, 0.8, "Industries we commission")
    counts, examples = {}, {}
    for st in data["stories"]:
        k = st["sector"]
        refs = [st["no"]] + st.get("merge", [])
        counts[k] = counts.get(k, 0) + len(refs)
        examples.setdefault(k, []).append(st.get("title") or projects[st["no"]]["name"]["en"])
    order = ["steam", "cc", "gt", "water", "oilgas", "solar", "industrial"]
    cw, ch, gap = 2.85, 2.25, 0.2
    for i, k in enumerate(order):
        r_, c_ = divmod(i, 4)
        x, y = 0.7 + c_ * (cw + gap), 1.95 + r_ * (ch + gap)
        rect(s, x, y, cw, ch, SAND)
        rect(s, x, y, cw, 0.07, ORANGE)
        text(s, x + 0.25, y + 0.25, 1.4, 0.8, str(counts.get(k, 0)), 40, H_XB, ORANGE)
        text(s, x + 0.25, y + 0.98, cw - 0.4, 0.4, data["sectors"][k], 13, H_B, NAVY, upper=True, spacing=0.5)
        ex = [e.split(" — ")[0].split(" (")[0] for e in examples.get(k, [])][:3]
        text(s, x + 0.25, y + 1.42, cw - 0.45, 0.8, ", ".join(ex), 10, B_R, SLATE, line=1.1)
    x, y = 0.7 + 3 * (cw + gap), 1.95 + (ch + gap)
    rect(s, x, y, cw, ch, NAVY)
    text(s, x + 0.25, y + 0.3, cw - 0.5, 1.8, [[("Counted by reference project. ", {"font": B_SB, "color": WHITE})],
         [("Most of our work is DCS commissioning on power plants — and the same procedures carry over to water, "
           "oil & gas and industrial sites.", {"color": "C9D3DC"})]], 11, B_R, WHITE, line=1.15, para_space=6)
    footer(s, n)


def where():
    s, n = new_slide()
    picture_fill(s, os.path.join(GEN, "maps", "overview.png"), 0, 0, 9.6, SH)
    rect(s, 9.6, 0, SW - 9.6, SH, NAVY)
    eyebrow(s, 9.95, 0.6, 3.2, "Where we have worked", ODARK)
    title(s, 9.95, 0.9, 3.2, 1.3, "Three countries, one team", size=30, color=WHITE)
    by = {}
    for p in projects.values():
        by[p["country"]] = by.get(p["country"], 0) + 1
    y = 2.6
    for c in ("egypt", "saudi-arabia", "oman"):
        text(s, 9.95, y, 1.2, 0.8, str(by.get(c, 0)), 40, H_XB, ODARK)
        text(s, 11.15, y + 0.12, 2.0, 0.6, [[(COUNTRY[c], {"font": H_B, "size": 17, "color": WHITE, "upper": True})],
             [("reference projects", {"size": 10, "color": "C9D3DC"})]], 10, B_R, WHITE)
        y += 1.15
    text(s, 9.95, 6.15, 3.1, 0.6, "Pins show plant locations; numbers are projects per site.", 9, B_R, "C9D3DC", line=1.1)
    footer(s, n, dark=True)
    notes(s, "Map: Natural Earth (public domain). Site locations are approximate (city level).")


def section(country_key, map_key, sites_text):
    s, n = new_slide()
    picture_fill(s, os.path.join(GEN, "maps", f"country-{map_key}.png"), 5.2, 0, SW - 5.2, SH)
    rect(s, 0, 0, 5.2, SH, NAVY)
    refs = [p for p in projects.values() if p["country"] == country_key]
    yrs = sorted([p["start"] for p in refs] + [p.get("end") or p["start"] for p in refs])
    eyebrow(s, 0.7, 1.6, 4, f"{yrs[0]}–{yrs[-1]}" if yrs[0] != yrs[-1] else str(yrs[0]), ODARK)
    title(s, 0.7, 1.95, 4.3, 1.4, COUNTRY[country_key], size=60, color=WHITE)
    text(s, 0.7, 3.2, 4.0, 0.6, f"{len(refs)} reference project{'s' if len(refs) > 1 else ''}", 20, H_B, ODARK, upper=True)
    text(s, 0.7, 3.85, 4.0, 2.8, sites_text, 12, B_R, "D5DEE6", line=1.2)
    footer(s, n, dark=True)


def logo_row(slide, keys, x, y, w, h=0.5):
    n = len(keys)
    gap = 0.15
    cell = min(1.7, (w - gap * (n - 1)) / n)
    for i, k in enumerate(keys):
        picture_fit(slide, logo_path(k), x + i * (cell + gap), y, cell, h, align="left")


def lines(t, per):
    """Rough line count for wrapped text at a given characters-per-line."""
    n, cur = 1, 0
    for w in t.split():
        if cur and cur + 1 + len(w) > per:
            n, cur = n + 1, len(w)
        else:
            cur += (1 if cur else 0) + len(w)
    return n


def story_slide(st):
    p = projects[st["no"]]
    s, n = new_slide()
    country, per = COUNTRY[p["country"]], period(p)
    plat = PLATFORM[p["platform"]]
    # photo column
    ref = st.get("photo") if photo_path(st.get("photo")) else st.get("fallbackPhoto")
    licensed = ref == st.get("photo") and ref and not ref.startswith(("pat:", "public:"))
    ph = photo_path(ref)
    px = 7.55
    if ph:
        picture_fill(s, ph, px, 0, SW - px, SH)
    else:
        rect(s, px, 0, SW - px, SH, SAND)
    rect(s, px, SH - 0.62, SW - px, 0.62, NAVY, alpha=0.7)
    if licensed and ref in credits:
        c = credits[ref]
        cap = f"Photo: {c.get('author') or 'Wikimedia Commons'} · {c.get('license', '')} · Wikimedia Commons"
    elif ref and ref.startswith("pat:") and st.get("photoCaption"):
        cap = st["photoCaption"]
    else:
        cap = "PAT on site — illustrative photo"
    text(s, px + 2.45, SH - 0.5, SW - px - 2.7, 0.4, cap, 8.5, B_R, "E5ECF2", anchor=MSO_ANCHOR.MIDDLE, line=1.0)
    # locator map inset
    loc = os.path.join(GEN, "maps", f"locator-{st['no']}.png")
    rect(s, px + 0.25, SH - 2.25, 2.06, 1.66, WHITE)
    picture_fill(s, loc, px + 0.28, SH - 2.22, 2.0, 1.6)
    # figure badge
    fig_v, fig_u = p["figure"]["value"], p["figure"]["unit"]["en"]
    rect(s, px + 0.25, 0.3, 2.0, 1.15, NAVY, alpha=0.88)
    text(s, px + 0.42, 0.36, 1.8, 0.7, fig_v, 30 if len(fig_v) <= 5 else 22, H_XB, ODARK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, px + 0.42, 1.03, 1.8, 0.3, fig_u, 9.5, H_SB, WHITE, upper=True, spacing=1.5)

    # text column
    L, W = 0.65, 6.55
    eyebrow(s, L, 0.5, W, " · ".join(x for x in (country, per, plat) if x))
    name = st.get("title") or p["name"]["en"]
    tsize = 28 if len(name) < 46 else 24
    tlines = lines(name, 40 if tsize == 28 else 47)
    title(s, L, 0.8, W, 1.25, name, size=tsize)
    ty = 0.8 + tlines * tsize / 60 + 0.2
    text(s, L, ty, W, 0.4, p["type"]["en"], 14, B_M, SLATE)

    # contractor / end user
    y0 = ty + 0.55
    rect(s, L, y0, W, 1.4, SAND)
    half = (W - 0.3) / 2
    text(s, L + 0.2, y0 + 0.14, half, 0.25, "Contractor", 9.5, H_SB, OTXT, spacing=1.8, upper=True)
    ckeys = st["contractors"]
    logo_row(s, ckeys, L + 0.2, y0 + 0.45, half - 0.15)
    text(s, L + 0.2, y0 + 1.02, half - 0.1, 0.35, " · ".join(companies[k]["name"] for k in ckeys), 9, B_M, NAVY, line=1.0)
    ex = L + 0.2 + half + 0.1
    rect(s, ex - 0.12, y0 + 0.18, 0.012, 1.05, LINE)
    text(s, ex, y0 + 0.14, half, 0.25, "End user", 9.5, H_SB, OTXT, spacing=1.8, upper=True)
    eu = st.get("endUser")
    if eu:
        logo_row(s, [eu], ex, y0 + 0.45, half - 0.3)
        text(s, ex, y0 + 1.02, half - 0.1, 0.35, companies[eu]["name"], 9, B_M, NAVY, line=1.0)
    else:
        text(s, ex, y0 + 0.5, half - 0.1, 0.7, "Owner to be confirmed", 11, B_M, SLATE)

    # description, scope, result
    y = y0 + 1.6
    text(s, L, y, W, 0.25, "The project", 9.5, H_SB, OTXT, spacing=1.8, upper=True)
    text(s, L, y + 0.27, W, 0.75, p["summary"]["en"], 11.5, B_R, NAVY, line=1.12)
    y += 0.27 + lines(p["summary"]["en"], 92) * 0.2 + 0.25
    text(s, L, y, W, 0.25, "PAT scope", 9.5, H_SB, OTXT, spacing=1.8, upper=True)
    if st.get("merge"):
        roles = st.get("roles", {})
        bullets = [f"{companies[k]['name'].split(' (')[0]} — {roles.get(k, '')}" for k in ckeys]
    else:
        bullets = [b["en"] for b in p["scope"]][:3]
    text(s, L, y + 0.27, W, 1.0, [[("•  ", {"color": ORANGE, "font": B_SB}), (b, {})] for b in bullets], 10.5, B_R, NAVY, line=1.05, para_space=2)
    y += 0.27 + sum(lines(b, 100) for b in bullets) * 0.18 + len(bullets) * 0.03 + 0.2
    if st.get("result"):
        ry = max(y, 6.0)
        rect(s, L, ry + 0.02, 0.06, 0.42, ORANGE)
        text(s, L + 0.2, ry, W - 0.2, 0.46, [[("Result  ", {"font": H_B, "color": OTXT, "upper": True, "spacing": 1.5}),
             (st["result"], {"font": B_M})]], 11, B_R, NAVY, anchor=MSO_ANCHOR.MIDDLE, line=1.0)
    footer(s, n)
    nt = [f"Reference no. {p['no']:02d}" + (f" (+ {', '.join(str(m) for m in st['merge'])})" if st.get("merge") else "") + ".",
          f"Contractor(s): {', '.join(companies[k]['name'] for k in ckeys)} — from the PAT reference list."]
    nt.append(f"End user: {companies[eu]['name'] if eu else 'not stated'}" + (" — TO CONFIRM (not in the reference list)." if st.get("confirmEndUser") else "."))
    nt.append("Photo: " + (f"{credits[ref]['source']} ({credits[ref].get('license', '')})" if licensed and ref in credits else "PAT photo" + (" (illustrative, not from this site)" if not st.get("photoCaption") else "")))
    if st.get("evidence"):
        nt.append(st["evidence"])
    notes(s, "\n".join(nt))


def recognition():
    s, n = new_slide()
    rect(s, 0, 0, 5.0, SH, SAND)
    picture_fit(s, os.path.join(ROOT, "public", "images", "emerson-plaque.jpg"), 0.6, 0.6, 3.8, 6.2)
    eyebrow(s, 5.6, 0.75, 7, "Recognition")
    title(s, 5.6, 1.05, 7.2, 1.3, "Recognised by Emerson Process Management")
    text(s, 5.6, 2.55, 7.0, 1.2, "“In recognition of Ahdy Abu-Malwa for technical contributions and excellence in delivery of Emerson "
         "Ovation Expert Control System Solutions within the Arab Republic of Egypt.”", 15, B_R, SLATE, line=1.2)
    text(s, 5.6, 3.95, 3.4, 2.2, [[("DCS Lead Engineer", {"font": H_B, "color": OTXT, "upper": True, "size": 12})],
         "New Talkha 750 MW CC Expansion", "Damietta 1,230 MW WDPF to Ovation Migration", "El-Kureimat II 750 MW CC Power Station"],
         11.5, B_R, NAVY, line=1.15, para_space=4)
    text(s, 9.2, 3.95, 3.4, 1.4, [[("DCS Engineer", {"font": H_B, "color": OTXT, "upper": True, "size": 12})],
         "Ataka 900 MW Steam Modernization"], 11.5, B_R, NAVY, line=1.15, para_space=4)
    text(s, 5.6, 6.2, 7.2, 0.5, "Signed by Robert Yeager, President, Emerson PWS, and T.C. Chan, Project Manager, Emerson PWS. "
         "Emerson and Ovation are trademarks of Emerson Electric Co.", 9, B_R, SLATE, line=1.1)
    footer(s, n)


def closing():
    s, n = new_slide()
    picture_fill(s, os.path.join(ROOT, "profile", "img", "team-night.jpg"), 0, 0, SW, SH)
    rect(s, 0, 0, SW, SH, NAVY, alpha=0.86)
    picture_fit(s, brand("png/logo/pat-weave-device-reverse-1600.png"), 7.6, 2.8, 7.2, 5.4)
    eyebrow(s, 0.75, 1.3, 8, "Your next outage or start-up", ODARK)
    title(s, 0.7, 1.65, 8.5, 1.9, "Let’s bring your control system online", size=48, color=WHITE)
    rows = [("Phone", "+20 109 991 2020  (24/7)"), ("E-mail", "info@patdcs.com  ·  sales@patdcs.com"), ("Web", "patdcs.com"),
            ("Office", "New Damietta, Egypt")]
    for i, (k, v) in enumerate(rows):
        text(s, 0.75, 3.85 + i * 0.5, 1.3, 0.4, k, 11, H_SB, ODARK, upper=True, spacing=1.5, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 2.05, 3.85 + i * 0.5, 6, 0.4, v, 15, B_M, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    qr = os.path.join(GEN, "qr-patdcs.png")
    q = qrcode.QRCode(border=2, box_size=12)
    q.add_data("https://patdcs.com")
    q.make()
    q.make_image(fill_color="#102840", back_color="white").save(qr)
    rect(s, 0.75, 6.0 - 0.02, 1.1, 1.1, WHITE)
    picture_fit(s, qr, 0.78, 6.0, 1.04, 1.04)
    text(s, 2.0, 6.25, 4, 0.6, "Scan for the full project list", 11, B_R, "D5DEE6")
    picture_fit(s, brand("png/logo/pat-lockup-horizontal-reverse-transparent-2000.png"), 0.75, 0.45, 3.0, 0.75, align="left")


# ------------------------------------------------------------------ deck
cover()
trusted_by()
numbers()
sectors()
where()
by_country = {"egypt": [], "saudi-arabia": [], "oman": []}
for st in data["stories"]:
    by_country[projects[st["no"]]["country"]].append(st)
SITES = {
    "egypt": "Damietta · Talkha · El Kureimat · Ataka (Suez) · Ain Sokhna · Al Arish · Alexandria\n\n"
             "Steam and combined-cycle power plants, solar ISCC, industrial and marine.",
    "saudi-arabia": "Shoaiba · Rabigh · Yanbu · Makkah · Riyadh · Al Qassim · Al Jouf · Al Qurayyat · Abqaiq\n\n"
                    "Steam and gas-turbine power plants, utilities, desalination and oil & gas processing.",
    "oman": "Sur\n\nA 2,000 MW combined-cycle IPP — Ovation DCS and Automatic Plant Start.",
}
for c, mk in (("egypt", "egypt"), ("saudi-arabia", "saudi"), ("oman", "oman")):
    section(c, mk, SITES[c])
    for st in sorted(by_country[c], key=lambda x: (projects[x["no"]]["start"], x["no"])):
        story_slide(st)
recognition()
closing()

out = os.path.join(OUT, "PAT-success-stories.pptx")
prs.save(out)
print("saved", out, len(prs.slides), "slides")
